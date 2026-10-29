#!/usr/bin/env python3
"""
Prospecta MEI - servidor local.

    python app.py                 # abre o sistema no navegador (http://localhost:8765)
    python app.py --sem-navegador
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import socket
import sqlite3
import subprocess
import sys
import threading
import time
import webbrowser
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

import atualizar_base as ab

PORTA = int(os.environ.get("PROSPECTA_PORTA", "8765"))
ESTATICOS = ab.BASE_DIR / "static"
ACOMP = ab.DADOS_DIR / "acompanhamento.db"
LOG_ATUALIZACAO = ab.DADOS_DIR / "atualizacao.log"
LIMITE_EXCEL = 200_000

SITUACOES = {1: "Nula", 2: "Ativa", 3: "Suspensa", 4: "Inapta", 8: "Baixada"}
STATUS = {"contatado": "Contatado", "interessado": "Interessado",
          "cliente": "Virou cliente", "descartado": "Sem interesse"}
PORTES = {"00": "Não informado", "01": "Microempresa", "03": "Empresa de pequeno porte", "05": "Demais"}
NATUREZAS = {"2135": "Empresário (individual)", "2062": "Sociedade empresária limitada",
             "2305": "Empresa individual de responsabilidade limitada", "2240": "Sociedade simples limitada"}

app = FastAPI(title="Prospecta MEI", docs_url=None, redoc_url=None)
app.mount("/static", StaticFiles(directory=ESTATICOS), name="static")


# --------------------------------------------------------------------------
# Banco
# --------------------------------------------------------------------------
_trava_troca = threading.Lock()


def preparar_acompanhamento() -> None:
    ab.DADOS_DIR.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(ACOMP) as db:
        db.execute("""CREATE TABLE IF NOT EXISTS contatos (
            cnpj TEXT PRIMARY KEY, status TEXT, nota TEXT, razao_social TEXT, atualizado_em TEXT)""")


def trocar_base_se_preciso() -> None:
    """Ativa a base nova deixada pela atualização (no Windows o arquivo só pode
    ser trocado quando nenhuma consulta está usando o antigo)."""
    if not (ab.MARCA_NOVO.exists() and ab.BANCO_NOVO.exists()):
        return
    with _trava_troca:
        if ab.MARCA_NOVO.exists() and ab.BANCO_NOVO.exists():
            try:
                os.replace(ab.BANCO_NOVO, ab.BANCO)
                ab.MARCA_NOVO.unlink(missing_ok=True)
            except PermissionError:
                pass


@contextmanager
def banco():
    trocar_base_se_preciso()
    if not ab.BANCO.exists():
        raise HTTPException(409, "A base ainda não foi baixada.")
    db = sqlite3.connect(f"file:{ab.BANCO.as_posix()}?mode=ro", uri=True)
    db.row_factory = sqlite3.Row
    db.execute("ATTACH DATABASE ? AS acomp", (str(ACOMP),))
    try:
        yield db
    finally:
        db.close()


def info_base() -> dict:
    trocar_base_se_preciso()
    return ab.ler_info_atual()


# --------------------------------------------------------------------------
# Formatação
# --------------------------------------------------------------------------
def fmt_cnpj(c: str) -> str:
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}" if c and len(c) == 14 else (c or "")


def fmt_tel(t: str | None) -> str:
    if not t:
        return ""
    ddd, n = (t[:2], t[2:]) if len(t) >= 10 else ("", t)
    corpo = f"{n[:-4]}-{n[-4:]}" if len(n) > 4 else n
    return f"({ddd}) {corpo}" if ddd else corpo


def whatsapp(t: str | None) -> str | None:
    if not ab.eh_celular(t):
        return None
    n = t[2:]
    return f"55{t[:2]}{'9' + n if len(n) == 8 else n}"


def fmt_data(iso: str | None) -> str:
    return f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}" if iso else ""


def fmt_cep(c: str | None) -> str:
    return f"{c[:5]}-{c[5:]}" if c and len(c) == 8 else (c or "")


def endereco(r) -> str:
    partes = [r["logradouro"] or ""]
    if r["numero"]:
        partes[0] += f", {r['numero']}"
    if r["complemento"]:
        partes.append(r["complemento"])
    return " - ".join(p for p in partes if p)


def item(r) -> dict:
    tels = [t for t in (r["telefone1"], r["telefone2"]) if t]
    return {
        "cnpj": r["cnpj"], "cnpj_fmt": fmt_cnpj(r["cnpj"]),
        "razao_social": r["razao_social"], "nome_fantasia": r["nome_fantasia"] or "",
        "tipo": r["tipo"], "situacao": SITUACOES.get(r["situacao"], str(r["situacao"])),
        "situacao_cod": r["situacao"], "data_inicio": r["data_inicio"],
        "data_exclusao_mei": r["data_exclusao_mei"],
        "cnae_principal": r["cnae_principal"], "cnae_descricao": r["cnae_descricao"] or "",
        "endereco": endereco(r), "bairro": r["bairro"] or "", "municipio": r["municipio"],
        "uf": r["uf"], "cep": fmt_cep(r["cep"]),
        "telefones": [{"numero": fmt_tel(t), "whatsapp": whatsapp(t)} for t in dict.fromkeys(tels)],
        "email": r["email"] or "", "status": r["status"], "nota": r["nota"] or "",
    }


# --------------------------------------------------------------------------
# Filtros
# --------------------------------------------------------------------------
DATA = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def montar_filtro(q) -> tuple[str, list, bool]:
    """Devolve (where, parâmetros, precisa_juntar_acompanhamento)."""
    cond: list[str] = []
    p: list = []

    def em(coluna: str, valores: list) -> None:
        cond.append(f"{coluna} IN ({','.join('?' * len(valores))})")
        p.extend(valores)

    if v := [u.upper() for u in q.getlist("uf") if u]:
        em("e.uf", v)
    if v := [m for m in q.getlist("municipio") if m]:
        em("e.municipio_codigo", v)
    if v := [ab.normalizar(b) for b in q.getlist("bairro") if b]:
        em("e.bairro_busca", v)
    if t := ab.normalizar(q.get("bairro_texto")):
        cond.append("e.bairro_busca LIKE ?")
        p.append(f"%{t}%")
    if cnaes := [re.sub(r"\D", "", c).zfill(7) for c in q.getlist("cnae") if re.sub(r"\D", "", c)]:
        if q.get("cnae_secundario") in ("1", "true"):
            partes = [f"e.cnae_principal IN ({','.join('?' * len(cnaes))})"]
            p.extend(cnaes)
            for c in cnaes:
                partes.append("e.cnaes_secundarios LIKE ?")
                p.append(f"%{c}%")
            cond.append("(" + " OR ".join(partes) + ")")
        else:
            em("e.cnae_principal", cnaes)
    if (tipo := q.get("tipo", "")) in ("MEI", "EX-MEI"):
        cond.append("e.tipo = ?")
        p.append(tipo)
    situacoes = [int(s) for s in q.getlist("situacao") if s.isdigit()]
    if situacoes:
        em("e.situacao", situacoes)
    for campo, coluna, op in (("aberto_de", "data_inicio", ">="), ("aberto_ate", "data_inicio", "<="),
                              ("excluido_de", "data_exclusao_mei", ">="),
                              ("excluido_ate", "data_exclusao_mei", "<=")):
        if (v := q.get(campo, "")) and DATA.match(v):
            cond.append(f"e.{coluna} {op} ?")
            p.append(v)
    contato = set(q.getlist("contato"))
    if "telefone" in contato:
        cond.append("(e.telefone1 IS NOT NULL OR e.telefone2 IS NOT NULL)")
    if "celular" in contato:
        cond.append("e.celular = 1")
    if "email" in contato:
        cond.append("e.email IS NOT NULL")
    if texto := (q.get("texto") or "").strip():
        compacto = re.sub(r"[^0-9A-Za-z]", "", texto).upper()
        if 8 <= len(compacto) <= 14 and sum(ch.isdigit() for ch in compacto) >= 6:
            cond.append("e.cnpj LIKE ?")
            p.append(compacto + "%")
        else:
            for palavra in ab.normalizar(texto).split()[:6]:
                cond.append("e.busca LIKE ?")
                p.append(f"%{palavra}%")
    junta = False
    if acomp := q.get("acompanhamento", ""):
        junta = True
        if acomp == "nao_contatado":
            cond.append("c.status IS NULL")
        elif acomp == "marcados":
            cond.append("c.status IS NOT NULL")
        elif acomp in STATUS:
            cond.append("c.status = ?")
            p.append(acomp)
    return (" WHERE " + " AND ".join(cond)) if cond else "", p, junta


ORDENS = {
    "recentes": "e.data_inicio DESC, e.cnpj",
    "antigos": "e.data_inicio ASC, e.cnpj",
    "nome": "e.razao_social COLLATE NOCASE, e.cnpj",
    "bairro": "e.municipio, e.bairro_busca, e.razao_social COLLATE NOCASE",
    "desenquadramento": "e.data_exclusao_mei IS NULL, e.data_exclusao_mei DESC, e.cnpj",
}
JUNCAO = " FROM empresas e LEFT JOIN acomp.contatos c ON c.cnpj = e.cnpj"
COLUNAS = "e.*, c.status, c.nota"


# --------------------------------------------------------------------------
# Rotas
# --------------------------------------------------------------------------
@app.get("/", include_in_schema=False)
def inicio():
    return FileResponse(ESTATICOS / "index.html", headers={"Cache-Control": "no-store"})


@app.get("/api/estado")
def estado():
    info = info_base()
    return {"base_pronta": bool(info), "info": info, "atualizacao": ler_progresso()}


@app.get("/api/opcoes")
def opcoes():
    with banco() as db:
        muns = db.execute("""SELECT municipio_codigo AS codigo, municipio AS nome, uf, COUNT(*) AS total
                             FROM empresas GROUP BY municipio_codigo ORDER BY nome""").fetchall()
        tipos = dict(db.execute("SELECT tipo, COUNT(*) FROM empresas GROUP BY tipo").fetchall())
    return {"municipios": [dict(m) for m in muns], "tipos": tipos}


@app.get("/api/bairros")
def bairros(request: Request):
    cods = [m for m in request.query_params.getlist("municipio") if m]
    if not cods:
        return []
    with banco() as db:
        linhas = db.execute(f"""SELECT bairro_busca AS chave, MIN(bairro) AS nome, COUNT(*) AS total
                                FROM empresas WHERE municipio_codigo IN ({','.join('?' * len(cods))})
                                  AND bairro_busca <> ''
                                GROUP BY bairro_busca ORDER BY bairro_busca LIMIT 2000""", cods).fetchall()
    return [dict(r) for r in linhas]


@app.get("/api/cnaes")
def cnaes(q: str = ""):
    cond, p = [], []
    for palavra in ab.normalizar(q).split()[:6]:
        cond.append("descricao_busca LIKE ?")
        p.append(f"%{palavra}%")
    where = (" WHERE " + " AND ".join(cond)) if cond else ""
    with banco() as db:
        linhas = db.execute(f"SELECT codigo, descricao, total FROM cnae_contagem{where} "
                            "ORDER BY total DESC LIMIT 30", p).fetchall()
    return [dict(r) for r in linhas]


@app.get("/api/buscar")
def buscar(request: Request):
    q = request.query_params
    where, p, junta = montar_filtro(q)
    ordem = ORDENS.get(q.get("ordem", ""), ORDENS["recentes"])
    pagina = max(1, int(q.get("pagina", "1") or 1))
    por_pagina = min(200, max(10, int(q.get("por_pagina", "50") or 50)))
    with banco() as db:
        origem = JUNCAO if junta else " FROM empresas e"
        total = db.execute(f"SELECT COUNT(*){origem}{where}", p).fetchone()[0]
        linhas = db.execute(f"SELECT {COLUNAS}{JUNCAO}{where} ORDER BY {ordem} LIMIT ? OFFSET ?",
                            [*p, por_pagina, (pagina - 1) * por_pagina]).fetchall()
    return {"total": total, "pagina": pagina, "por_pagina": por_pagina, "itens": [item(r) for r in linhas]}


@app.get("/api/resumo")
def resumo(request: Request):
    where, p, junta = montar_filtro(request.query_params)
    origem = JUNCAO if junta else " FROM empresas e"
    with banco() as db:
        def top(expr, rotulo):
            return [dict(r) for r in db.execute(
                f"SELECT {expr} AS chave, {rotulo} AS nome, COUNT(*) AS total{origem}{where} "
                f"GROUP BY {expr} ORDER BY total DESC LIMIT 8", p).fetchall()]
        contagens = db.execute(f"""SELECT COUNT(*) AS total, SUM(e.celular) AS celular,
                                   SUM(e.email IS NOT NULL) AS email{origem}{where}""", p).fetchone()
        return {
            "municipios": top("e.municipio_codigo", "MIN(e.municipio)"),
            "bairros": top("e.municipio_codigo || '|' || e.bairro_busca", "MIN(e.bairro) || ' (' || MIN(e.municipio) || ')'"),
            "atividades": top("e.cnae_principal", "MIN(e.cnae_descricao)"),
            "contagens": dict(contagens),
        }


@app.get("/api/empresa/{cnpj}")
def empresa(cnpj: str):
    with banco() as db:
        r = db.execute(f"SELECT {COLUNAS}{JUNCAO} WHERE e.cnpj = ?", (cnpj.upper(),)).fetchone()
        if not r:
            raise HTTPException(404, "Empresa não encontrada na base.")
        secs = (r["cnaes_secundarios"] or "").split(",") if r["cnaes_secundarios"] else []
        desc = dict(db.execute(f"SELECT codigo, descricao FROM cnaes WHERE codigo IN ({','.join('?' * len(secs))})",
                               secs).fetchall()) if secs else {}
        atualizado = db.execute("SELECT atualizado_em FROM acomp.contatos WHERE cnpj = ?", (r["cnpj"],)).fetchone()
    dados = item(r)
    dados.update({
        "secundarias": [{"codigo": c, "descricao": desc.get(c, "")} for c in secs],
        "data_situacao": r["data_situacao"], "data_opcao_mei": r["data_opcao_mei"],
        "opcao_simples": r["opcao_simples"], "data_exclusao_simples": r["data_exclusao_simples"],
        "natureza": NATUREZAS.get(r["natureza_juridica"] or "", r["natureza_juridica"] or ""),
        "porte": PORTES.get(r["porte"] or "", r["porte"] or ""), "capital_social": r["capital_social"],
        "acompanhamento_atualizado": atualizado[0] if atualizado else None,
    })
    return dados


class Acompanhamento(BaseModel):
    status: str | None = None
    nota: str = ""
    razao_social: str = ""


@app.put("/api/acompanhamento/{cnpj}")
def salvar_acompanhamento(cnpj: str, dados: Acompanhamento):
    status = dados.status if dados.status in STATUS else None
    nota = dados.nota.strip()[:4000]
    with sqlite3.connect(ACOMP) as db:
        if not status and not nota:
            db.execute("DELETE FROM contatos WHERE cnpj = ?", (cnpj.upper(),))
        else:
            db.execute("""INSERT INTO contatos VALUES (?,?,?,?,?)
                          ON CONFLICT(cnpj) DO UPDATE SET status=excluded.status, nota=excluded.nota,
                          razao_social=excluded.razao_social, atualizado_em=excluded.atualizado_em""",
                       (cnpj.upper(), status, nota, dados.razao_social, datetime.now().isoformat(timespec="seconds")))
    return {"ok": True, "status": status, "nota": nota}


# --------------------------------------------------------------------------
# Exportação
# --------------------------------------------------------------------------
CAMPOS_EXPORT = [
    ("CNPJ", 20), ("Razão social", 38), ("Nome fantasia", 28), ("Tipo", 9), ("Situação", 10),
    ("Abertura", 12), ("Saiu do MEI em", 14), ("CNAE", 10), ("Atividade principal", 45),
    ("Endereço", 40), ("Bairro", 22), ("Município", 20), ("UF", 5), ("CEP", 11),
    ("Telefone 1", 16), ("Telefone 2", 16), ("WhatsApp", 32), ("E-mail", 32),
    ("Status do contato", 16), ("Anotações", 40),
]


def linha_export(r) -> list:
    tels = list(dict.fromkeys(t for t in (r["telefone1"], r["telefone2"]) if t))
    zap = next((whatsapp(t) for t in tels if whatsapp(t)), None)
    return [
        fmt_cnpj(r["cnpj"]), r["razao_social"], r["nome_fantasia"] or "", r["tipo"],
        SITUACOES.get(r["situacao"], ""), r["data_inicio"], r["data_exclusao_mei"],
        r["cnae_principal"] or "", r["cnae_descricao"] or "", endereco(r), r["bairro"] or "",
        r["municipio"] or "", r["uf"], fmt_cep(r["cep"]),
        fmt_tel(tels[0]) if tels else "", fmt_tel(tels[1]) if len(tels) > 1 else "",
        f"https://wa.me/{zap}" if zap else "", r["email"] or "",
        STATUS.get(r["status"] or "", ""), r["nota"] or "",
    ]


def nome_arquivo(q, ext: str) -> str:
    partes = ["ex-meis" if q.get("tipo") == "EX-MEI" else "meis"]
    cods = q.getlist("municipio")
    if len(cods) == 1:
        with banco() as db:
            r = db.execute("SELECT municipio FROM empresas WHERE municipio_codigo = ? LIMIT 1", (cods[0],)).fetchone()
        if r:
            partes.append(r[0])
    if len(bairros := q.getlist("bairro")) == 1:
        partes.append(bairros[0])
    partes.append(date.today().strftime("%Y-%m-%d"))
    nome = "-".join(re.sub(r"[^a-z0-9]+", "-", ab.normalizar(p).lower()).strip("-") for p in partes)
    return f"{nome}.{ext}"


@app.get("/api/exportar")
def exportar(request: Request):
    q = request.query_params
    formato = q.get("formato", "xlsx")
    where, p, _ = montar_filtro(q)
    ordem = ORDENS.get(q.get("ordem", ""), ORDENS["recentes"])
    sql = f"SELECT {COLUNAS}{JUNCAO}{where} ORDER BY {ordem}"

    if formato == "csv":
        def gerar():
            with banco() as db:
                buf = io.StringIO()
                w = csv.writer(buf, delimiter=";")
                buf.write("﻿")
                w.writerow([c for c, _ in CAMPOS_EXPORT])
                for n, r in enumerate(db.execute(sql, p), 1):
                    linha = linha_export(r)
                    linha[5], linha[6] = fmt_data(linha[5]), fmt_data(linha[6])
                    w.writerow(linha)
                    if n % 2000 == 0:
                        yield buf.getvalue()
                        buf.seek(0)
                        buf.truncate()
                yield buf.getvalue()
        return StreamingResponse(gerar(), media_type="text/csv; charset=utf-8",
                                 headers={"Content-Disposition": f'attachment; filename="{nome_arquivo(q, "csv")}"'})

    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Empresas")
    ws.freeze_panes = "C2"
    for i, (_, larg) in enumerate(CAMPOS_EXPORT, 1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    cab_fonte, cab_fundo = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F5C45")
    cabecalho = []
    for titulo, _ in CAMPOS_EXPORT:
        c = WriteOnlyCell(ws, value=titulo)
        c.font, c.fill, c.alignment = cab_fonte, cab_fundo, Alignment(vertical="center")
        cabecalho.append(c)
    ws.append(cabecalho)
    n = 0
    with banco() as db:
        for r in db.execute(sql + " LIMIT ?", [*p, LIMITE_EXCEL]):
            linha = linha_export(r)
            for k in (5, 6):
                if linha[k]:
                    cel = WriteOnlyCell(ws, value=date.fromisoformat(linha[k]))
                    cel.number_format = "DD/MM/YYYY"
                    linha[k] = cel
            if linha[16]:
                cel = WriteOnlyCell(ws, value=linha[16])
                cel.hyperlink = linha[16]
                cel.font = Font(color="1F5C45", underline="single")
                linha[16] = cel
            ws.append(linha)
            n += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CAMPOS_EXPORT))}{n + 1}"
    saida = io.BytesIO()
    wb.save(saida)
    return Response(saida.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nome_arquivo(q, "xlsx")}"'})


# --------------------------------------------------------------------------
# Atualização da base pela interface
# --------------------------------------------------------------------------
def ler_progresso() -> dict | None:
    try:
        estado = json.loads(ab.PROGRESSO.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if estado.get("status") == "executando":
        ultima = datetime.fromisoformat(estado.get("atualizado_em", "2000-01-01"))
        if (datetime.now() - ultima).total_seconds() > 900:
            estado["status"] = "erro"
            estado["mensagem"] = "A atualização parou de responder. Clique em atualizar para continuar de onde parou."
    return estado


@app.post("/api/atualizacao")
def iniciar_atualizacao(forcar: bool = False):
    if ab.outra_atualizacao_rodando():
        return {"ok": False, "mensagem": "Já existe uma atualização em andamento."}
    ufs = info_base().get("ufs") or os.environ.get("PROSPECTA_UF", "CE")
    ab.DADOS_DIR.mkdir(parents=True, exist_ok=True)
    ab.PROGRESSO.write_text(json.dumps({
        "status": "executando", "etapa": "Iniciando", "mensagem": "Conectando à Receita Federal...",
        "detalhe": "", "percentual": 0, "iniciado_em": datetime.now().isoformat(timespec="seconds"),
        "atualizado_em": datetime.now().isoformat(timespec="seconds"), "log": [],
    }, ensure_ascii=False), encoding="utf-8")
    cmd = [sys.executable, str(ab.BASE_DIR / "atualizar_base.py"), "--uf", ufs]
    if forcar:
        cmd.append("--forcar")
    extra = {"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {"start_new_session": True}
    with open(LOG_ATUALIZACAO, "w", encoding="utf-8") as log:
        subprocess.Popen(cmd, stdout=log, stderr=subprocess.STDOUT, cwd=ab.BASE_DIR,
                         env={**os.environ, "PYTHONIOENCODING": "utf-8"}, **extra)
    return {"ok": True}


@app.get("/api/atualizacao")
def progresso():
    return ler_progresso() or {"status": "parado"}


# --------------------------------------------------------------------------
def porta_ocupada(porta: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.settimeout(0.5)
        return s.connect_ex(("127.0.0.1", porta)) == 0


def main() -> None:
    import uvicorn

    endereco_local = f"http://localhost:{PORTA}"
    abrir = "--sem-navegador" not in sys.argv
    if porta_ocupada(PORTA):
        print(f"O sistema já está aberto em {endereco_local}")
        if abrir:
            webbrowser.open(endereco_local)
        return
    preparar_acompanhamento()
    if abrir:
        threading.Thread(target=lambda: (time.sleep(1.5), webbrowser.open(endereco_local)), daemon=True).start()
    print(f"\n  Prospecta MEI rodando em {endereco_local}\n  Para fechar, feche esta janela.\n")
    uvicorn.run(app, host="127.0.0.1", port=PORTA, log_level="warning")


preparar_acompanhamento()

if __name__ == "__main__":
    main()
